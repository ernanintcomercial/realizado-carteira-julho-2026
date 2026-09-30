"""Aplica uma vez a transferência aprovada da Santil (Walter -> Daniel).

Usa as linhas atribuídas ao Daniel no PD019 e preserva todos os totais.
Executar somente sobre a publicação anterior à transferência.
"""
import argparse
import json
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

MONTHS = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
          "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]
CLIENTS = {10964, 921, 17106, 17671, 18069, 22089, 22580}


def apply(payload, amounts):
    september = next(p for p in payload["periodos"] if p["id"] == "M09")
    before = sum(Decimal(str(r["meta"])) for r in september["registros"] if r["repId"] == 9963)
    if before != Decimal("583988.93"):
        raise ValueError(f"Meta-base do Daniel diferente da aprovada: {before}; não reaplicar a transferência.")
    for period in payload["periodos"]:
        transfer = sum((amounts.get(month, Decimal(0)) for month in period["meses"]), Decimal(0))
        if not transfer:
            continue
        source = next(r for r in period["registros"] if r["repId"] == 4803 and r["contrato"] == "PLÁSTICO")
        target = next(r for r in period["registros"] if r["repId"] == 9963 and r["contrato"] == "PLÁSTICO")
        total_before = sum(Decimal(str(r["meta"])) for r in period["registros"])
        if Decimal(str(source["meta"])) < transfer:
            raise ValueError(f"Saldo de meta insuficiente no Walter: {period['id']}")
        source["meta"] = float(Decimal(str(source["meta"])) - transfer)
        target["meta"] = float(Decimal(str(target["meta"])) + transfer)
        assert total_before == sum(Decimal(str(r["meta"])) for r in period["registros"])
        # O mês corrente também possui uma meta proporcional ao corte.
        for record in (source, target):
            if "metaAteCorte" in record:
                daily = payload["metaDiaria"]
                ratio = Decimal(str(daily["diasUteisTranscorridos"])) / Decimal(str(daily["diasUteisMes"]))
                record["metaAteCorte"] = float((Decimal(str(record["meta"])) * ratio).quantize(Decimal("0.01")))
    after = sum(Decimal(str(r["meta"])) for r in september["registros"] if r["repId"] == 9963)
    assert after == Decimal("639933.32"), after
    return payload


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--pd019", type=Path, required=True)
    parser.add_argument("--dashboard", type=Path, required=True)
    args = parser.parse_args()
    wb = load_workbook(args.pd019, read_only=True, data_only=True)
    sheet = wb.active
    headers = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
    amounts = {month: Decimal(0) for month in range(1, 13)}
    seen = set()
    for row in sheet.iter_rows(min_row=3, values_only=True):
        if row[0] != 2026 or row[1] != 9963 or row[3] not in CLIENTS:
            continue
        if str(row[7]).upper() != "DP" or str(row[9]).upper() != "VENDAS":
            continue
        if row[3] in seen:
            raise ValueError(f"Cliente duplicado na meta do Daniel: {row[3]}")
        seen.add(row[3])
        for month, name in enumerate(MONTHS, 1):
            amounts[month] += Decimal(str(row[headers.index(name)] or 0))
    wb.close()
    assert seen == CLIENTS, seen
    assert amounts[9] == Decimal("55944.39"), amounts
    payload = json.loads(args.dashboard.read_text(encoding="utf-8"))
    updated = apply(payload, amounts)
    args.dashboard.write_text(json.dumps(updated, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print("Transferência aplicada; meta setembro Daniel 639933.32; totais preservados.")


if __name__ == "__main__":
    main()
